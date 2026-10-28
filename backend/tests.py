"""Decision-logic tests (no models or photos needed).

    python tests.py
"""
import unittest

import numpy as np

from config import Config
from fusion import Fusion
from gallery import Gallery

rng = np.random.default_rng(42)


def unit(v):
    return (v / np.linalg.norm(v, axis=-1, keepdims=True)).astype(np.float32)


def school(n=300, photos=20, noise=0.03):
    centers = unit(rng.normal(size=(n, 512)))
    centers[1] = unit(centers[0] * 0.85 + centers[1] * 0.15)       # 1 is a look-alike of 0
    per = {}
    for i in range(n):
        photos_i = unit(centers[i] + rng.normal(scale=noise, size=(photos, 512)))
        far = unit(photos_i + rng.normal(scale=noise * 2, size=(photos, 512)))       # stand-in for the far versions
        per[str(1000 + i)] = [np.stack([p, f]) for p, f in zip(photos_i, far)]
    g = Gallery.build(per, {})
    g.calibrate(3, 0.55)
    return g, centers


class Decisions(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.g, cls.c = school()

    def probe(self, i, noise=0.03):
        return unit(self.c[i] + rng.normal(scale=noise, size=512))

    def run_track(self, frames, live=0.95, cfg=None, width=120):
        f = Fusion(cfg or Config())
        r = None
        for emb in frames:
            r = f.update("t", emb, 1.0, width, live, self.g)
            if r["status"] != "pending":
                break
        return r

    def test_enrolled_student_is_matched_after_min_frames(self):
        r = self.run_track([self.probe(5) for _ in range(10)], cfg=Config(fast_track=False))
        self.assertEqual(r["status"], "match")
        self.assertEqual(r["student_id"], "1005")
        self.assertEqual(r["frames"], 3)

    def test_single_frame_never_marks_without_fast_track(self):
        f = Fusion(Config(fast_track=False))
        self.assertEqual(f.update("t", self.probe(5), 1.0, 120, 0.95, self.g)["status"], "pending")

    def test_fast_track_overwhelming_match_in_one_frame(self):
        r = Fusion(Config()).update("t", self.probe(6, noise=0.01), 1.0, 150, 0.97, self.g)
        self.assertEqual((r["status"], r["student_id"], r["frames"]), ("match", "1006", 1))

    def test_fast_track_never_without_liveness_or_for_far_faces(self):
        self.assertEqual(Fusion(Config()).update("t", self.probe(7, noise=0.01), 1.0, 150, 0.6, self.g)["status"], "pending")
        self.assertEqual(Fusion(Config()).update("t", self.probe(7, noise=0.01), 1.0, 150, None, self.g)["status"], "pending")
        self.assertEqual(Fusion(Config()).update("t", self.probe(7, noise=0.01), 1.0, 40, 0.97, self.g)["status"], "pending")

    def test_fast_track_never_for_lookalikes(self):
        r = Fusion(Config()).update("t", self.probe(0, noise=0.01), 1.0, 150, 0.97, self.g)
        self.assertEqual(r["status"], "pending")

    def test_stranger_is_unknown(self):
        stranger = unit(rng.normal(size=512))
        r = self.run_track([unit(stranger + rng.normal(scale=0.03, size=512)) for _ in range(10)])
        self.assertEqual(r["status"], "unknown")

    def test_lookalikes_are_flagged(self):
        self.assertIn("1001", self.g.lookalikes.get("1000", []))

    def test_ambiguous_between_lookalikes_goes_to_review(self):
        mid = unit(self.c[0] + self.c[1])
        r = self.run_track([unit(mid + rng.normal(scale=0.01, size=512)) for _ in range(20)])
        self.assertEqual(r["status"], "review")

    def test_photo_attack_is_spoof(self):
        r = self.run_track([self.probe(7) for _ in range(10)], live=0.1)
        self.assertEqual(r["status"], "spoof")

    def test_threshold_is_calibrated_and_sane(self):
        self.assertEqual(self.g.calib["source"], "calibrated")
        for near in (True, False):
            th = self.g.thresholds(near)
            self.assertGreater(th["t_accept"], th["t_review"])
        self.assertGreater(self.g.probability(0.7), 0.99)
        self.assertLess(self.g.probability(0.1), 0.01)

    def test_save_and_load_round_trip(self):
        import tempfile
        with tempfile.TemporaryDirectory() as d:
            self.g.save(d)
            g2 = Gallery.load(d)
        q = self.probe(9)
        self.assertEqual(self.g.match(q)[0][0], g2.match(q)[0][0])
        self.assertAlmostEqual(self.g.match(q)[0][1], g2.match(q)[0][1], places=2)


    def test_far_face_is_recognised_but_waits_for_liveness(self):
        r = self.run_track([self.probe(11) for _ in range(30)], live=None, width=36)
        self.assertEqual(r["status"], "pending")
        self.assertEqual(r["student_id"], "1011")
        self.assertEqual(r["range"], "far")
        self.assertIn("closer", r["reason"])

    def test_far_then_near_is_matched(self):
        f = Fusion(Config())
        for _ in range(6):
            r = f.update("t", self.probe(12), 0.3, 36, None, self.g)
        for _ in range(3):
            r = f.update("t", self.probe(12), 1.0, 90, 0.95, self.g)
            if r["status"] != "pending":
                break
        self.assertEqual((r["status"], r["student_id"]), ("match", "1012"))

    def test_identity_switch_restarts_the_track(self):
        f = Fusion(Config(fast_track=False))
        f.update("t", self.probe(20), 1.0, 120, 0.95, self.g)
        f.update("t", self.probe(20), 1.0, 120, 0.95, self.g)
        r = f.update("t", self.probe(21), 1.0, 120, 0.95, self.g)      # tracker swapped people
        self.assertEqual(r["frames"], 1)
        for _ in range(3):
            r = f.update("t", self.probe(21), 1.0, 120, 0.95, self.g)
        self.assertEqual((r["status"], r["student_id"]), ("match", "1021"))

    def test_four_students_at_once(self):
        f = Fusion(Config())
        done = {}
        for _ in range(4):
            for i in (30, 31, 32, 33):
                r = f.update(f"t{i}", self.probe(i), 1.0, 100, 0.95, self.g)
                if r["status"] == "match":
                    done[f"t{i}"] = r["student_id"]
        self.assertEqual(done, {"t30": "1030", "t31": "1031", "t32": "1032", "t33": "1033"})

    def test_confident_match_teaches_the_gallery(self):
        f = Fusion(Config())
        learned = None
        for _ in range(3):
            r = f.update("t", self.probe(40, noise=0.01), 1.0, 120, 0.95, self.g)
            learned = r.get("_learn", learned)
        self.assertEqual(r["status"], "match")
        self.assertIsNotNone(learned)
        g = self.g
        g.add_adaptive("1040", learned, 5)
        self.assertEqual(int((g.adapt_owner == g.index["1040"]).sum()), 1)
        for _ in range(7):
            g.add_adaptive("1040", self.probe(40), 5)
        self.assertEqual(int((g.adapt_owner == g.index["1040"]).sum()), 5)     # only the newest few are kept
        self.assertEqual(g.match(self.probe(40))[0][0], "1040")


class IncrementalEnrolment(unittest.TestCase):
    def test_add_replace_keep_learning(self):
        g, c = school(n=50)
        g.add_adaptive("1003", unit(c[3] + rng.normal(scale=0.02, size=512)), 5)
        newcomer = unit(rng.normal(size=512))
        photos = [np.stack([p, p]) for p in unit(newcomer + rng.normal(scale=0.03, size=(20, 512)))]
        g2 = g.with_students({"2000": photos}, {"2000": "2000/a.jpg"}, 0.55)
        self.assertEqual(len(g2.ids), 51)
        self.assertEqual(g2.match(unit(newcomer + rng.normal(scale=0.03, size=512)))[0][0], "2000")
        self.assertEqual(g2.match(unit(c[7] + rng.normal(scale=0.03, size=512)))[0][0], "1007")   # others untouched
        self.assertEqual(int((g2.adapt_owner == g2.index["1003"]).sum()), 1)                     # gate learning kept
        self.assertEqual(g2.calib, g.calib)
        # a re-take replaces the student's photos instead of adding to them
        g3 = g2.with_students({"2000": photos[:8]}, {}, 0.55)
        self.assertEqual(int(g3.sp_count[g3.index["2000"]]), 8)
        # a twin enrolled later is linked as a look-alike both ways
        twin = [np.stack([p, p]) for p in unit(c[5] * 0.9 + rng.normal(scale=0.03, size=(20, 512)))]
        g4 = g3.with_students({"2001": twin}, {}, 0.55)
        self.assertIn("1005", g4.lookalikes.get("2001", []))
        self.assertIn("2001", g4.lookalikes.get("1005", []))


class Tracking(unittest.TestCase):
    def test_four_people_keep_their_ids(self):
        from tracker import Tracker
        t = Tracker()
        boxes = [np.array([x, 100, x + 60, 170], np.float32) for x in (100, 400, 700, 1000)]
        first = t.assign(boxes)
        self.assertEqual(len(set(first)), 4)
        moved = [b + np.array([15, 10, 18, 14], np.float32) for b in boxes]     # everyone walks a bit closer
        self.assertEqual(t.assign(moved[::-1]), first[::-1])

    def test_new_person_gets_new_id(self):
        from tracker import Tracker
        t = Tracker()
        a = t.assign([np.array([100, 100, 160, 170], np.float32)])
        b = t.assign([np.array([100, 100, 160, 170], np.float32), np.array([900, 100, 940, 150], np.float32)])
        self.assertEqual(b[0], a[0])
        self.assertNotEqual(b[1], a[0])


class DailyReport(unittest.TestCase):
    def setUp(self):
        import tempfile
        from records import Store
        from report import Reporter
        self.tmp = tempfile.TemporaryDirectory()
        self.store = Store(self.tmp.name)
        roster = [{"id": str(i), "name": f"S{i}", "cls": c, "section": "A", "roll": str(i)} for i, c in enumerate(["X", "X", "IX", "XII"], 1)]
        self.store.put_roster(roster, False, "h")
        import datetime as dt
        self.today = dt.date.today().isoformat()
        self.store.put_day(self.today, {"1": {"s": "P", "t": "07:50", "src": "cam", "c": 0.99}, "2": {"s": "L", "t": "08:20", "src": "cam"},
                                        "3": {"s": "A", "t": "09:15", "src": "finalize"}}, "d")
        self.r = Reporter(self.tmp.name, self.tmp.name + "/reports", self.store)
        self.sent = []
        import report
        self._send = report.send
        report.send = lambda settings, msg: self.sent.append(msg)

    def tearDown(self):
        import report
        report.send = self._send
        self.tmp.cleanup()

    def configure(self, **kw):
        self.r.update_settings({"enabled": True, "recipients": ["a@b.org"], "time": "00:00", "days": list(range(7)),
                                "smtp": {"host": "smtp.x", "port": 587, "username": "me@x.org", "password": "pw"}, **kw})

    def test_workbook_and_summary(self):
        import io
        import openpyxl
        from report import build_workbook, summarise
        roster, _ = self.store.roster()
        s = summarise(roster, self.store.day(self.today))
        self.assertEqual((s["total"]["P"], s["total"]["L"], s["total"]["A"], s["total"]["unmarked"]), (1, 1, 1, 1))
        self.assertEqual([c["cls"] for c in s["classes"]], ["IX-A", "X-A", "XII-A"])      # school order, not alphabetical
        wb = openpyxl.load_workbook(io.BytesIO(build_workbook(self.today, roster, self.store.day(self.today))))
        self.assertEqual(wb.sheetnames[:3], ["Summary", "Absentees", "Late"])
        absentees = [r[3] for r in wb["Absentees"].iter_rows(min_row=5, values_only=True)]
        self.assertEqual(sorted(absentees), ["S3", "S4"])                                  # absent + not seen

    def test_not_ready_without_settings(self):
        with self.assertRaises(ValueError):
            self.r.run(self.today)

    def test_scheduler_sends_once_a_day(self):
        self.configure()
        self.r.tick(); self.r.tick()
        self.assertEqual(len(self.sent), 1)
        self.assertIn("2/4 present", self.sent[0]["Subject"])

    def test_manual_send_does_not_cancel_the_scheduled_one(self):
        self.configure()
        self.r.run(self.today)
        self.r.tick()
        self.assertEqual(len(self.sent), 2)

    def test_skips_days_without_attendance(self):
        self.configure()
        import os
        os.remove(os.path.join(self.store.days_dir, f"{self.today}.json"))
        self.r.tick()
        self.assertEqual(self.sent, [])

    def test_password_is_never_returned(self):
        self.configure()
        self.assertNotIn("password", self.r.public_settings()["smtp"])
        self.r.update_settings({"smtp": {"password": ""}})                                  # blank keeps the old one
        self.assertEqual(self.r.settings["smtp"]["password"], "pw")


if __name__ == "__main__":
    unittest.main(verbosity=2)
