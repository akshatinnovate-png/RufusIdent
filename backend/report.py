"""Daily attendance report: an Excel workbook emailed on school days.

Attendance reaches the backend from the dashboard (see records.py). At the set
time the backend builds the workbook, emails it to the recipients and keeps a
copy in reports_dir (point it at OneDrive). If the laptop was off at that
time, the report goes out as soon as the backend starts again that day.
"""
import datetime as dt
import io
import json
import os
import smtplib
import ssl
import threading
import time
from email.message import EmailMessage
from email.utils import formataddr, formatdate, make_msgid

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SCHOOL = "Gulmohur High School"
STATUS = {"P": "Present", "L": "Late", "A": "Absent", None: "Not seen"}
SOURCE = {"cam": "Gate camera", "gate-manual": "Gate (manual)", "manual": "Register", "assistant": "AI assistant",
          "finalize": "Marked absent", "demo": "Demo"}
DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]

DEFAULT_SETTINGS = {
    "enabled": False,
    "recipients": [],
    "time": "09:30",
    "days": [0, 1, 2, 3, 4, 5],            # Monday..Saturday
    "smtp": {"host": "smtp.gmail.com", "port": 587, "security": "starttls", "username": "", "password": "",
             "from_name": "RufusIdent · Gulmohur High School"},
    "reports_dir": "",
}

GREEN, AMBER, RED, NAVY, GREY = "E6F4EA", "FFF4D6", "FDE6DC", "1D2160", "F3F4F6"
THIN = Side(style="thin", color="E5E7EB")

ROMAN = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5, "VI": 6, "VII": 7, "VIII": 8, "IX": 9, "X": 10, "XI": 11, "XII": 12}


def class_key(s: dict) -> str:
    return f"{s['cls']}-{s['section']}"


def class_order(key: str):
    cls, _, sec = key.partition("-")
    c = cls.strip().upper()
    return (ROMAN.get(c) or (int(c) if c.isdigit() else 99), sec)


def roll_order(s: dict):
    try:
        return (0, int(s.get("roll") or 0), s["name"])
    except ValueError:
        return (1, 0, s["name"])


# ───────────────────────── numbers ─────────────────────────

def summarise(roster: list[dict], records: dict) -> dict:
    classes: dict[str, dict] = {}
    for s in roster:
        c = classes.setdefault(class_key(s), {"cls": class_key(s), "total": 0, "P": 0, "L": 0, "A": 0, "unmarked": 0, "students": []})
        c["total"] += 1
        r = records.get(s["id"])
        c[r["s"] if r else "unmarked"] += 1
        c["students"].append(s)
    rows = sorted(classes.values(), key=lambda c: class_order(c["cls"]))
    for c in rows:
        c["here"] = c["P"] + c["L"]
        c["pct"] = round(c["here"] / c["total"] * 100, 1) if c["total"] else 0.0
        c["students"].sort(key=roll_order)
    tot = {k: sum(c[k] for c in rows) for k in ("total", "P", "L", "A", "unmarked")}
    tot["here"] = tot["P"] + tot["L"]
    tot["pct"] = round(tot["here"] / tot["total"] * 100, 1) if tot["total"] else 0.0
    return {"classes": rows, "total": tot}


# ───────────────────────── workbook ─────────────────────────

def _header(ws, row: int, labels: list[str], widths: list[int]):
    for i, (label, w) in enumerate(zip(labels, widths), 1):
        cell = ws.cell(row=row, column=i, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 20


def _title(ws, text: str, sub: str):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14, color="0B3A20")
    ws["A2"] = sub
    ws["A2"].font = Font(italic=True, color="6B7280")


def _status_fill(status: str | None) -> PatternFill | None:
    return {"P": PatternFill("solid", fgColor=GREEN), "L": PatternFill("solid", fgColor=AMBER),
            "A": PatternFill("solid", fgColor=RED), None: PatternFill("solid", fgColor=GREY)}.get(status)


def build_workbook(date: str, roster: list[dict], records: dict, demo: bool = False) -> bytes:
    summary = summarise(roster, records)
    nice = dt.date.fromisoformat(date).strftime("%A, %d %B %Y")
    made = dt.datetime.now().strftime("%d %b %Y, %H:%M")
    wb = Workbook()

    # Summary
    ws = wb.active
    ws.title = "Summary"
    _title(ws, f"{SCHOOL} · Attendance · {nice}" + (" · DEMO DATA" if demo else ""), f"Generated {made} by RufusIdent. 'Not seen' = not marked by the gate or a teacher.")
    _header(ws, 4, ["Class", "Strength", "Present", "Late", "Absent", "Not seen", "Attendance %"], [12, 11, 10, 8, 9, 10, 14])
    r = 5
    for c in summary["classes"]:
        for i, v in enumerate([c["cls"], c["total"], c["P"], c["L"], c["A"], c["unmarked"], c["pct"] / 100], 1):
            cell = ws.cell(row=r, column=i, value=v)
            cell.border = Border(bottom=THIN)
        ws.cell(row=r, column=7).number_format = "0.0%"
        r += 1
    t = summary["total"]
    for i, v in enumerate(["TOTAL", t["total"], t["P"], t["L"], t["A"], t["unmarked"], t["pct"] / 100], 1):
        cell = ws.cell(row=r, column=i, value=v)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=GREY)
    ws.cell(row=r, column=7).number_format = "0.0%"
    if r > 5:
        rng = f"G5:G{r - 1}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0.75"], fill=PatternFill("solid", fgColor=RED)))
        ws.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["0.75", "0.9"], fill=PatternFill("solid", fgColor=AMBER)))
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["0.9"], fill=PatternFill("solid", fgColor=GREEN)))
    ws.freeze_panes = "A5"

    # Absent + not seen: the list the office acts on.
    ws = wb.create_sheet("Absentees")
    _title(ws, f"Absent or not seen · {nice}", "Sorted by class and roll number.")
    _header(ws, 4, ["Class", "Roll", "Admission No", "Name", "Status"], [10, 7, 15, 30, 12])
    r = 5
    for c in summary["classes"]:
        for s in c["students"]:
            rec = records.get(s["id"])
            if rec and rec["s"] != "A":
                continue
            st = rec["s"] if rec else None
            for i, v in enumerate([c["cls"], _num(s.get("roll")), s["id"], s["name"], STATUS[st]], 1):
                ws.cell(row=r, column=i, value=v).fill = _status_fill(st)
            r += 1
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:E{max(r - 1, 4)}"

    # Late arrivals
    ws = wb.create_sheet("Late")
    _title(ws, f"Late arrivals · {nice}", "Arrived after the late cut-off.")
    _header(ws, 4, ["Class", "Roll", "Admission No", "Name", "Time"], [10, 7, 15, 30, 10])
    r = 5
    late = [(c["cls"], s, records[s["id"]]) for c in summary["classes"] for s in c["students"] if records.get(s["id"], {}).get("s") == "L"]
    for cls, s, rec in sorted(late, key=lambda x: x[2].get("t", "")):
        for i, v in enumerate([cls, _num(s.get("roll")), s["id"], s["name"], rec.get("t", "")], 1):
            ws.cell(row=r, column=i, value=v)
        r += 1
    ws.freeze_panes = "A5"

    # One sheet per class
    for c in summary["classes"]:
        ws = wb.create_sheet(_sheet_name(c["cls"]))
        _title(ws, f"Class {c['cls']} · {nice}", f"{c['here']} of {c['total']} present ({c['pct']}%) · {c['L']} late · {c['A']} absent · {c['unmarked']} not seen")
        _header(ws, 4, ["Roll", "Admission No", "Name", "Status", "Time", "Marked by", "Match confidence"], [7, 15, 30, 11, 8, 16, 17])
        for r, s in enumerate(c["students"], 5):
            rec = records.get(s["id"])
            st = rec["s"] if rec else None
            vals = [_num(s.get("roll")), s["id"], s["name"], STATUS[st], (rec or {}).get("t", ""),
                    SOURCE.get((rec or {}).get("src"), (rec or {}).get("src", "")), (rec or {}).get("c", "")]
            for i, v in enumerate(vals, 1):
                cell = ws.cell(row=r, column=i, value=v)
                cell.border = Border(bottom=THIN)
            ws.cell(row=r, column=4).fill = _status_fill(st)
            if isinstance(vals[6], float):
                ws.cell(row=r, column=7).number_format = "0.0%"
        ws.freeze_panes = "A5"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _num(v):
    try:
        return int(v)
    except (TypeError, ValueError):
        return v or ""


def _sheet_name(name: str) -> str:
    for ch in '\\/?*[]:':
        name = name.replace(ch, "")
    return name[:31] or "Class"


# ───────────────────────── email ─────────────────────────

def compose(settings: dict, date: str, roster: list[dict], records: dict, demo: bool, xlsx: bytes) -> EmailMessage:
    s = summarise(roster, records)
    t = s["total"]
    day = dt.date.fromisoformat(date)
    nice = day.strftime("%a %d %b %Y")
    msg = EmailMessage()
    msg["Subject"] = f"{'[DEMO DATA] ' if demo else ''}GHS Attendance · {nice} · {t['here']:,}/{t['total']:,} present ({t['pct']}%)"
    smtp = settings["smtp"]
    msg["From"] = formataddr((smtp.get("from_name") or "RufusIdent", smtp["username"]))
    msg["To"] = ", ".join(settings["recipients"])
    # A real Date and a Message-ID on the sender's own domain: mail without them is
    # often quarantined silently by school / Microsoft 365 mail filters.
    msg["Date"] = formatdate(localtime=True)
    msg["Message-ID"] = make_msgid(domain=smtp["username"].rpartition("@")[2] or None)

    low = [c for c in s["classes"] if c["pct"] < 85]
    text = [f"{SCHOOL} — attendance for {day.strftime('%A, %d %B %Y')}", "",
            f"Present: {t['P']:,}   Late: {t['L']:,}   Absent: {t['A']:,}   Not seen: {t['unmarked']:,}   of {t['total']:,} ({t['pct']}%)", ""]
    if low:
        text.append("Classes below 85%: " + ", ".join(f"{c['cls']} ({c['pct']}%)" for c in low))
    text += ["", "The full class-wise sheet is attached.", "", "— RufusIdent · GHS AI Team"]
    msg.set_content("\n".join(text))

    rows = "".join(
        f"<tr><td style='padding:4px 10px'>{c['cls']}</td><td style='padding:4px 10px;text-align:right'>{c['here']}/{c['total']}</td>"
        f"<td style='padding:4px 10px;text-align:right'>{c['L']}</td><td style='padding:4px 10px;text-align:right'>{c['A'] + c['unmarked']}</td>"
        f"<td style='padding:4px 10px;text-align:right;font-weight:600;color:{'#9a2f12' if c['pct'] < 75 else '#8a6200' if c['pct'] < 90 else '#136b35'}'>{c['pct']}%</td></tr>"
        for c in s["classes"])
    html = f"""<div style="font-family:Segoe UI,Arial,sans-serif;color:#111827;max-width:640px">
  <div style="background:#0b3a20;color:#fff;padding:16px 20px;border-radius:12px 12px 0 0">
    <div style="font-size:12px;letter-spacing:.12em;opacity:.8">GULMOHUR HIGH SCHOOL · DAILY ATTENDANCE</div>
    <div style="font-size:22px;font-weight:700;margin-top:4px">{day.strftime('%A, %d %B %Y')}{' · DEMO DATA' if demo else ''}</div>
  </div>
  <div style="border:1px solid #e5e7eb;border-top:0;padding:16px 20px;border-radius:0 0 12px 12px">
    <p style="font-size:30px;font-weight:700;margin:0">{t['pct']}% <span style="font-size:15px;color:#6b7280;font-weight:400">present ({t['here']:,} of {t['total']:,})</span></p>
    <p style="margin:8px 0 16px;color:#374151">Present {t['P']:,} · Late {t['L']:,} · Absent {t['A']:,} · Not seen {t['unmarked']:,}</p>
    {"<p style='background:#fde6dc;color:#9a2f12;padding:8px 12px;border-radius:8px'>Below 85%: " + ", ".join(f"<b>{c['cls']}</b> {c['pct']}%" for c in low) + "</p>" if low else ""}
    <table style="border-collapse:collapse;font-size:14px;width:100%">
      <tr style="background:#f3f4f6;text-align:left"><th style="padding:6px 10px">Class</th><th style="padding:6px 10px;text-align:right">Present</th>
      <th style="padding:6px 10px;text-align:right">Late</th><th style="padding:6px 10px;text-align:right">Absent / not seen</th><th style="padding:6px 10px;text-align:right">%</th></tr>
      {rows}
    </table>
    <p style="color:#6b7280;font-size:13px;margin-top:16px">The attached Excel file has the summary, the absentee list, late arrivals and one sheet per class.</p>
    <p style="color:#9ca3af;font-size:12px">RufusIdent · Made by GHS AI Team · Developed by Akshat Sarkar</p>
  </div>
</div>"""
    msg.add_alternative(html, subtype="html")
    msg.add_attachment(xlsx, maintype="application", subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       filename=f"GHS_Attendance_{date}.xlsx")
    return msg


def send(settings: dict, msg: EmailMessage):
    smtp = settings["smtp"]
    password = os.environ.get("RUFUS_SMTP_PASSWORD") or smtp.get("password", "")
    host, port, security = smtp["host"], int(smtp["port"]), smtp.get("security", "starttls")
    ctx = ssl.create_default_context()
    if security == "ssl":
        server = smtplib.SMTP_SSL(host, port, context=ctx, timeout=30)
    else:
        server = smtplib.SMTP(host, port, timeout=30)
    with server:
        if security == "starttls":
            server.starttls(context=ctx)
        if smtp.get("username") and password:
            server.login(smtp["username"], password)
        refused = server.send_message(msg)
    return refused or {}


def friendly_error(e: Exception) -> str:
    if isinstance(e, smtplib.SMTPAuthenticationError):
        return ("The mail server rejected the username/password. For Gmail use an App Password "
                "(Google Account → Security → 2-Step Verification → App passwords), not the normal password.")
    if isinstance(e, (TimeoutError, ConnectionError, OSError)) and not isinstance(e, smtplib.SMTPException):
        return f"Could not reach the mail server ({e}). Check the internet connection and the server/port."
    return str(e) or e.__class__.__name__


# ───────────────────────── settings, state, scheduler ─────────────────────────

class Reporter:
    def __init__(self, data_dir: str, default_reports_dir: str, store):
        self.data_dir = data_dir
        self.default_reports_dir = default_reports_dir
        self.store = store
        self.settings_path = os.path.join(data_dir, "report_settings.json")
        self.state_path = os.path.join(data_dir, "report_state.json")
        self.lock = threading.Lock()
        self.settings = self._read(self.settings_path, DEFAULT_SETTINGS)
        self.settings["smtp"] = {**DEFAULT_SETTINGS["smtp"], **self.settings.get("smtp", {})}
        self.state = self._read(self.state_path, {"last_sent_date": None, "last_error": None, "last_attempt": 0, "history": []})

    @staticmethod
    def _read(path: str, default: dict) -> dict:
        try:
            with open(path, encoding="utf-8") as f:
                return {**json.loads(json.dumps(default)), **json.load(f)}
        except (FileNotFoundError, json.JSONDecodeError):
            return json.loads(json.dumps(default))

    def _write(self, path: str, obj: dict):
        os.makedirs(self.data_dir, exist_ok=True)
        tmp = path + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(obj, f, indent=1)
        os.replace(tmp, path)

    def public_settings(self) -> dict:
        s = json.loads(json.dumps(self.settings))
        s["smtp"]["has_password"] = bool(s["smtp"].pop("password", "") or os.environ.get("RUFUS_SMTP_PASSWORD"))
        return s

    def update_settings(self, patch: dict) -> dict:
        with self.lock:
            smtp = patch.pop("smtp", {}) or {}
            if not smtp.get("password"):
                smtp.pop("password", None)          # blank = keep the saved one
            self.settings.update({k: v for k, v in patch.items() if k in DEFAULT_SETTINGS})
            self.settings["smtp"].update({k: v for k, v in smtp.items() if k in DEFAULT_SETTINGS["smtp"]})
            self.settings["recipients"] = [r.strip() for r in self.settings["recipients"] if r.strip()]
            self._write(self.settings_path, self.settings)
        return self.public_settings()

    def status(self) -> dict:
        nxt = self.next_run()
        return {**{k: v for k, v in self.state.items() if k != "history"}, "history": self.state["history"][-10:],
                "next_run": nxt.isoformat(timespec="minutes") if nxt else None,
                "ready": self._problem() is None, "problem": self._problem()}

    def _problem(self) -> str | None:
        s = self.settings
        if not s["recipients"]:
            return "No recipients set."
        if not s["smtp"].get("host") or not s["smtp"].get("username"):
            return "Mail account not set up."
        if not (s["smtp"].get("password") or os.environ.get("RUFUS_SMTP_PASSWORD")):
            return "Mail password not set."
        return None

    def next_run(self) -> dt.datetime | None:
        if not self.settings["enabled"] or not self.settings["days"]:
            return None
        hh, mm = map(int, self.settings["time"].split(":"))
        now = dt.datetime.now()
        for add in range(8):
            day = (now + dt.timedelta(days=add)).date()
            at = dt.datetime.combine(day, dt.time(hh, mm))
            if day.weekday() in self.settings["days"] and (add or self.state["last_sent_date"] != day.isoformat()):
                return max(at, now) if add == 0 else at
        return None

    def build(self, date: str) -> tuple[bytes, list, dict, bool]:
        roster, demo = self.store.roster()
        records = self.store.day(date)
        if not roster:
            raise ValueError("No student roster has reached the backend yet. Open the dashboard once so it can sync.")
        return build_workbook(date, roster, records, demo), roster, records, demo

    def save_copy(self, date: str, xlsx: bytes) -> str:
        folder = self.settings.get("reports_dir") or self.default_reports_dir
        os.makedirs(folder, exist_ok=True)
        path = os.path.join(folder, f"GHS_Attendance_{date}.xlsx")
        with open(path, "wb") as f:
            f.write(xlsx)
        return path

    def run(self, date: str | None = None, test: bool = False, auto: bool = False) -> dict:
        """Build and email the report. Returns what happened."""
        date = date or dt.date.today().isoformat()
        with self.lock:
            problem = self._problem()
            if problem:
                raise ValueError(problem)
            xlsx, roster, records, demo = self.build(date)
            copy = self.save_copy(date, xlsx)
            msg = compose(self.settings, date, roster, records, demo, xlsx)
            if test:
                msg.replace_header("Subject", "[TEST] " + msg["Subject"])
            self.state["last_attempt"] = time.time()
            try:
                refused = send(self.settings, msg) or {}
            except Exception as e:
                self.state["last_error"] = friendly_error(e)
                self._log(date, False, self.state["last_error"], auto, test)
                raise RuntimeError(self.state["last_error"]) from e
            accepted = [r for r in self.settings["recipients"] if r not in refused]
            if not accepted:
                self.state["last_error"] = "The mail server refused every recipient: " + ", ".join(refused)
                self._log(date, False, self.state["last_error"], auto, test)
                raise RuntimeError(self.state["last_error"])
            self.state["last_error"] = None
            if auto:
                self.state["last_sent_date"] = date      # a manual send never cancels the scheduled one
            detail = f"accepted by {self.settings['smtp']['host']} for {', '.join(accepted)}"
            if refused:
                detail += f" · REFUSED: {', '.join(refused)}"
            self._log(date, True, detail, auto, test)
            return {"ok": True, "date": date, "recipients": accepted, "refused": list(refused), "saved": copy,
                    "subject": msg["Subject"], "message_id": msg["Message-ID"]}

    def _log(self, date, ok, detail, auto, test):
        self.state["history"].append({"at": dt.datetime.now().isoformat(timespec="seconds"), "date": date, "ok": ok,
                                      "detail": detail, "auto": auto, "test": test})
        self.state["history"] = self.state["history"][-60:]
        self._write(self.state_path, self.state)

    def tick(self):
        """Called every 30 s: send today's report once, at or after the set time."""
        s = self.settings
        now = dt.datetime.now()
        today = now.date().isoformat()
        if not s["enabled"] or now.weekday() not in s["days"] or self.state["last_sent_date"] == today:
            return
        if now.strftime("%H:%M") < s["time"] or self._problem():
            return
        if not self.store.day(today):
            return            # nothing marked today (holiday?): wait in case marks still arrive
        if self.state["last_error"] and time.time() - self.state["last_attempt"] < 600:
            return            # failed recently: retry every 10 minutes
        try:
            self.run(today, auto=True)
        except Exception:
            pass              # recorded in state; retried on a later tick

    def start(self):
        def loop():
            while True:
                try:
                    self.tick()
                except Exception:
                    pass
                time.sleep(30)
        threading.Thread(target=loop, name="daily-report", daemon=True).start()
